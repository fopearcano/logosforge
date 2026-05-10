"""Export project data to JSON, Markdown, CSV, DOCX, Fountain, PDF, HTML, or FDX."""

import csv
import io
import json
import xml.etree.ElementTree as ET

from storyplanner.db import Database


def _gather_project_data(db: Database, project_id: int) -> dict:
    project = db.get_project_by_id(project_id)

    characters = db.get_all_characters(project_id)
    places = db.get_all_places(project_id)
    notes = db.get_all_notes(project_id)
    scenes = db.get_all_scenes(project_id)

    char_name_by_id = {c.id: c.name for c in characters}
    place_name_by_id = {p.id: p.name for p in places}

    scene_list = []
    for i, scene in enumerate(scenes):
        char_ids = db.get_scene_character_ids(scene.id)
        place_ids = db.get_scene_place_ids(scene.id)
        char_states = db.get_scene_character_states(scene.id)

        scene_list.append(
            {
                "title": scene.title,
                "summary": scene.summary,
                "synopsis": scene.synopsis,
                "goal": scene.goal,
                "conflict": scene.conflict,
                "outcome": scene.outcome,
                "content": scene.content,
                "beat": scene.beat,
                "tags": [t.strip() for t in scene.tags.split(",") if t.strip()] if scene.tags else [],
                "order_index": i + 1,
                "act": scene.act,
                "chapter": scene.chapter,
                "plotline": scene.plotline,
                "characters": [
                    char_name_by_id[cid]
                    for cid in char_ids
                    if cid in char_name_by_id
                ],
                "places": [
                    place_name_by_id[pid]
                    for pid in place_ids
                    if pid in place_name_by_id
                ],
                "character_states": [
                    {"character": char_name_by_id[cid], "state": state}
                    for cid, state in char_states
                    if cid in char_name_by_id
                ],
            }
        )

    psyke_entries = db.get_all_psyke_entries(project_id)
    psyke_name_by_id = {e.id: e.name for e in psyke_entries}
    psyke_names_lower = {e.name.lower() for e in psyke_entries}
    characters = [c for c in characters if c.name.lower() not in psyke_names_lower]
    places = [p for p in places if p.name.lower() not in psyke_names_lower]

    scene_title_by_id = {s.id: s.title for s in scenes}

    psyke_list = []
    for e in psyke_entries:
        related = db.get_related_psyke_entries(e.id)
        progressions = db.get_psyke_progressions(e.id)
        psyke_list.append({
            "name": e.name,
            "entry_type": e.entry_type,
            "aliases": e.aliases,
            "notes": e.notes,
            "is_global": e.is_global,
            "details": db.get_psyke_entry_details(e.id),
            "related_entries": [r.name for r in related],
            "progressions": [
                {
                    "text": p.text,
                    "scene_title": scene_title_by_id.get(p.scene_id, "")
                    if p.scene_id
                    else "",
                    "sort_order": p.sort_order,
                }
                for p in progressions
            ],
        })

    outline_nodes = db.get_outline_nodes(project_id)
    children_map: dict[int | None, list] = {}
    for node in outline_nodes:
        children_map.setdefault(node.parent_id, []).append(node)

    def _build_outline_tree(parent_id: int | None) -> list[dict]:
        children = children_map.get(parent_id, [])
        children.sort(key=lambda n: (n.sort_order, n.id))
        return [
            {
                "title": n.title,
                "description": n.description,
                "children": _build_outline_tree(n.id),
            }
            for n in children
        ]

    from storyplanner.quantum_outliner.persistence import export_quantum_state

    data = {
        "project": {
            "title": project.title if project else "Untitled",
            "description": project.description if project else "",
            "format_mode": (project.format_mode if project else "novel") or "novel",
        },
        "characters": [
            {"name": c.name, "description": c.description} for c in characters
        ],
        "places": [
            {"name": p.name, "description": p.description} for p in places
        ],
        "notes": [
            {
                "title": n.title,
                "content": n.content,
                "tags": n.tags,
                "pinned": n.pinned,
                "psyke_links": [
                    psyke_name_by_id[eid]
                    for eid in db.get_note_psyke_links(n.id)
                    if eid in psyke_name_by_id
                ],
                "scene_links": [
                    scene_title_by_id[sid]
                    for sid in db.get_note_scene_links(n.id)
                    if sid in scene_title_by_id
                ],
            }
            for n in notes
        ],
        "scenes": scene_list,
        "psyke_entries": psyke_list,
        "outline": _build_outline_tree(None),
    }

    quantum = export_quantum_state(db, project_id)
    if quantum is not None:
        data["quantum_state"] = quantum

    return data


def export_json(db: Database, project_id: int) -> str:
    data = _gather_project_data(db, project_id)
    return json.dumps(data, indent=2, ensure_ascii=False)


def export_markdown(db: Database, project_id: int) -> str:
    data = _gather_project_data(db, project_id)
    lines: list[str] = []

    # Project header
    lines.append(f"# {data['project']['title']}")
    if data["project"]["description"]:
        lines.append("")
        lines.append(data["project"]["description"])

    # Characters
    lines.append("")
    lines.append("## Characters")
    if data["characters"]:
        for char in data["characters"]:
            lines.append("")
            lines.append(f"### {char['name']}")
            if char["description"]:
                lines.append("")
                lines.append(char["description"])
    else:
        lines.append("")
        lines.append("No characters.")

    # Places
    lines.append("")
    lines.append("## Places")
    if data["places"]:
        for place in data["places"]:
            lines.append("")
            lines.append(f"### {place['name']}")
            if place["description"]:
                lines.append("")
                lines.append(place["description"])
    else:
        lines.append("")
        lines.append("No places.")

    # Notes
    lines.append("")
    lines.append("## Notes")
    if data["notes"]:
        for note in data["notes"]:
            lines.append("")
            lines.append(f"### {note['title']}")
            if note["content"]:
                lines.append("")
                lines.append(note["content"])
    else:
        lines.append("")
        lines.append("No notes.")

    # Scenes
    lines.append("")
    lines.append("## Scenes")
    if data["scenes"]:
        for scene in data["scenes"]:
            lines.append("")
            lines.append(f"### {scene['order_index']}. {scene['title']}")
            if scene["act"]:
                lines.append(f"- **Act:** {scene['act']}")
            if scene["chapter"]:
                lines.append(f"- **Chapter:** {scene['chapter']}")
            if scene["plotline"]:
                lines.append(f"- **Plotline:** {scene['plotline']}")
            if scene["beat"]:
                lines.append(f"- **Beat:** {scene['beat']}")
            if scene["tags"]:
                lines.append(f"- **Tags:** {', '.join(scene['tags'])}")
            if scene["characters"]:
                lines.append(f"- **Characters:** {', '.join(scene['characters'])}")
            if scene["places"]:
                lines.append(f"- **Places:** {', '.join(scene['places'])}")
            if scene["summary"]:
                lines.append("")
                lines.append(scene["summary"])
            if scene["synopsis"]:
                lines.append("")
                lines.append(f"**Synopsis:** {scene['synopsis']}")
            if scene["goal"]:
                lines.append(f"- **Goal:** {scene['goal']}")
            if scene["conflict"]:
                lines.append(f"- **Conflict:** {scene['conflict']}")
            if scene["outcome"]:
                lines.append(f"- **Outcome:** {scene['outcome']}")
            if scene["character_states"]:
                lines.append("")
                lines.append("**Character States:**")
                for cs in scene["character_states"]:
                    lines.append(f"- {cs['character']}: {cs['state']}")
            if scene["content"]:
                lines.append("")
                lines.append("---")
                lines.append("")
                lines.append(scene["content"])
    else:
        lines.append("")
        lines.append("No scenes.")

    lines.append("")
    return "\n".join(lines)


def export_outline_markdown(db: Database, project_id: int) -> str:
    data = _gather_project_data(db, project_id)
    lines: list[str] = []

    lines.append(f"# {data['project']['title']}")

    scenes = data["scenes"]
    if not scenes:
        lines.append("")
        lines.append("No scenes to display.")
        lines.append("")
        return "\n".join(lines)

    chapter_groups = _group_scenes_by_chapter(scenes)

    for chapter_name, group_scenes in chapter_groups:
        lines.append("")
        if chapter_name:
            lines.append(f"## {chapter_name}")
        else:
            lines.append("## Uncategorized")

        for scene in group_scenes:
            lines.append("")
            lines.append(f"### {scene['order_index']}. {scene['title']}")
            if scene["act"]:
                lines.append(f"- **Act:** {scene['act']}")
            if scene["plotline"]:
                lines.append(f"- **Plotline:** {scene['plotline']}")
            if scene["beat"]:
                lines.append(f"- **Beat:** {scene['beat']}")
            if scene["tags"]:
                lines.append(f"- **Tags:** {', '.join(scene['tags'])}")
            if scene["characters"]:
                lines.append(f"- **Characters:** {', '.join(scene['characters'])}")
            if scene["places"]:
                lines.append(f"- **Places:** {', '.join(scene['places'])}")
            if scene["summary"]:
                lines.append("")
                lines.append(scene["summary"])
            if scene["synopsis"]:
                lines.append("")
                lines.append(f"**Synopsis:** {scene['synopsis']}")
            if scene["goal"]:
                lines.append(f"- **Goal:** {scene['goal']}")
            if scene["conflict"]:
                lines.append(f"- **Conflict:** {scene['conflict']}")
            if scene["outcome"]:
                lines.append(f"- **Outcome:** {scene['outcome']}")
            if scene["character_states"]:
                lines.append("")
                lines.append("**Character States:**")
                for cs in scene["character_states"]:
                    lines.append(f"- {cs['character']}: {cs['state']}")

    lines.append("")
    return "\n".join(lines)


def export_csv_scenes(db: Database, project_id: int) -> str:
    data = _gather_project_data(db, project_id)

    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(
        ["order_index", "title", "summary", "synopsis", "goal", "conflict", "outcome",
         "beat", "tags", "act", "chapter", "plotline", "characters", "places"]
    )
    for scene in data["scenes"]:
        writer.writerow(
            [
                scene["order_index"],
                scene["title"],
                scene["summary"],
                scene["synopsis"],
                scene["goal"],
                scene["conflict"],
                scene["outcome"],
                scene["beat"],
                ", ".join(scene["tags"]),
                scene["act"],
                scene["chapter"],
                scene["plotline"],
                ", ".join(scene["characters"]),
                ", ".join(scene["places"]),
            ]
        )
    return output.getvalue()


def _group_scenes_by_chapter(scenes: list[dict]) -> list[tuple[str, list[dict]]]:
    groups: list[tuple[str, list[dict]]] = []
    current_chapter: str | None = None
    current_group: list[dict] = []

    for scene in scenes:
        chapter = scene["chapter"] if scene["chapter"] else ""
        if chapter != current_chapter:
            if current_group:
                groups.append((current_chapter or "", current_group))
            current_chapter = chapter
            current_group = [scene]
        else:
            current_group.append(scene)
    if current_group:
        groups.append((current_chapter or "", current_group))

    return groups


def _scene_body(scene: dict) -> str:
    return scene["content"] or scene["synopsis"] or scene["summary"] or ""


def export_screenplay(db: Database, project_id: int) -> str:
    data = _gather_project_data(db, project_id)
    return _format_text(data, "screenplay")


def export_manuscript(db: Database, project_id: int) -> str:
    data = _gather_project_data(db, project_id)
    return _format_text(data, "novel")


def export_formatted_text(db: Database, project_id: int) -> str:
    data = _gather_project_data(db, project_id)
    return _format_text(data, _get_fmt(data))


def _is_script_format(fmt: str) -> bool:
    return fmt in ("screenplay", "series", "stage_script", "graphic_novel")


def _format_text(data: dict, fmt: str) -> str:
    if fmt in ("screenplay", "series"):
        return _fmt_screenplay_text(data, fmt)
    if fmt == "stage_script":
        return _fmt_stage_script_text(data)
    if fmt == "graphic_novel":
        return _fmt_graphic_novel_text(data)
    return _fmt_novel_text(data)


def _slug_line(scene: dict) -> str:
    if scene["places"]:
        place_str = ", ".join(scene["places"]).upper()
        return f"INT. {place_str} — {scene['title'].upper()}"
    return scene["title"].upper()


def _fmt_novel_text(data: dict) -> str:
    lines: list[str] = []
    title = data["project"]["title"]
    lines.append(title)
    lines.append("=" * len(title))
    lines.append("")

    if not data["scenes"]:
        lines.append("No scenes.")
        return "\n".join(lines)

    chapter_groups = _group_scenes_by_chapter(data["scenes"])

    for chapter_name, group_scenes in chapter_groups:
        if chapter_name:
            lines.append("")
            lines.append(chapter_name)
            lines.append("-" * len(chapter_name))
            lines.append("")

        for scene in group_scenes:
            lines.append(scene["title"])
            lines.append("")
            body = _scene_body(scene)
            if body:
                lines.append(body)
            lines.append("")
            lines.append("")

    return "\n".join(lines)


def _fmt_screenplay_text(data: dict, fmt: str) -> str:
    lines: list[str] = []
    title = data["project"]["title"]
    lines.append(title.upper())
    lines.append("=" * len(title))
    lines.append("")
    lines.append("")

    current_act = None
    for scene in data["scenes"]:
        act = scene.get("act", "")
        if fmt == "series" and act and act != current_act:
            current_act = act
            lines.append("")
            lines.append(f"        {act.upper()}")
            lines.append("")

        lines.append(_slug_line(scene))
        lines.append("")
        body = _scene_body(scene)
        if body:
            lines.append(body)
            lines.append("")
        lines.append("")

    return "\n".join(lines)


def _fmt_stage_script_text(data: dict) -> str:
    lines: list[str] = []
    title = data["project"]["title"]
    lines.append(f"        {title.upper()}")
    lines.append("")
    lines.append("")

    current_act = None
    current_scene_num = 0
    for scene in data["scenes"]:
        act = scene.get("act", "")
        if act and act != current_act:
            current_act = act
            current_scene_num = 0
            lines.append("")
            lines.append(f"        {act.upper()}")
            lines.append("")

        current_scene_num += 1
        lines.append(f"        SCENE {current_scene_num}")
        lines.append("")

        body = _scene_body(scene)
        if body:
            lines.append(body)
            lines.append("")
        lines.append("")

    return "\n".join(lines)


def _fmt_graphic_novel_text(data: dict) -> str:
    lines: list[str] = []
    title = data["project"]["title"]
    lines.append(title.upper())
    lines.append("")

    page_num = 0
    for scene in data["scenes"]:
        page_num += 1
        lines.append(f"PAGE {page_num}")
        lines.append("")
        body = _scene_body(scene)
        if body:
            lines.append(body)
            lines.append("")
        lines.append("")

    return "\n".join(lines)


# -- DOCX export (format-aware) -----------------------------------------------

_SCRIPT_FONT_DOCX = "Courier New"
_PROSE_FONT_DOCX = "Times New Roman"


def _get_fmt(data: dict) -> str:
    return data["project"].get("format_mode", "novel")


def export_docx_manuscript(db: Database, project_id: int, path: str) -> None:
    data = _gather_project_data(db, project_id)
    fmt = _get_fmt(data)
    if fmt in ("screenplay", "series"):
        _docx_screenplay(data, path, fmt)
    elif fmt == "stage_script":
        _docx_stage_script(data, path)
    elif fmt == "graphic_novel":
        _docx_graphic_novel(data, path)
    else:
        _docx_novel(data, path)


def _docx_title_page(doc, title: str, font_name: str) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    title_para = doc.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_para.add_run(title)
    title_run.bold = True
    title_run.font.size = Pt(24)
    title_run.font.name = font_name
    title_para.paragraph_format.space_after = Pt(0)
    doc.add_page_break()


def _docx_novel(data: dict, path: str) -> None:
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = _PROSE_FONT_DOCX
    style.font.size = Pt(12)
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.line_spacing = 1.15

    _docx_title_page(doc, data["project"]["title"], _PROSE_FONT_DOCX)

    if not data["scenes"]:
        doc.add_paragraph("No scenes.")
        doc.save(path)
        return

    chapter_groups = _group_scenes_by_chapter(data["scenes"])
    chapter_num = 0
    for chapter_name, group_scenes in chapter_groups:
        if chapter_name:
            chapter_num += 1
            heading = f"Chapter {chapter_num}: {chapter_name}"
        else:
            heading = "Scenes"
        ch_para = doc.add_paragraph()
        ch_para.paragraph_format.space_before = Pt(24)
        ch_run = ch_para.add_run(heading)
        ch_run.bold = True
        ch_run.font.size = Pt(16)
        ch_run.font.name = _PROSE_FONT_DOCX

        for scene in group_scenes:
            scene_para = doc.add_paragraph()
            scene_para.paragraph_format.space_before = Pt(12)
            scene_run = scene_para.add_run(scene["title"])
            scene_run.italic = True
            scene_run.font.size = Pt(12)
            scene_run.font.name = _PROSE_FONT_DOCX

            body = _scene_body(scene)
            if body:
                _add_content_paragraphs(doc, body, _PROSE_FONT_DOCX)

    doc.save(path)


def _docx_screenplay(data: dict, path: str, fmt: str) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = _SCRIPT_FONT_DOCX
    style.font.size = Pt(12)
    style.paragraph_format.space_after = Pt(0)
    style.paragraph_format.line_spacing = 1.0

    for section in doc.sections:
        section.left_margin = Inches(1.5)
        section.right_margin = Inches(1.0)
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)

    _docx_title_page(doc, data["project"]["title"], _SCRIPT_FONT_DOCX)

    if not data["scenes"]:
        doc.add_paragraph("No scenes.")
        doc.save(path)
        return

    current_act = None
    for scene in data["scenes"]:
        act = scene.get("act", "")
        if fmt == "series" and act and act != current_act:
            current_act = act
            act_para = doc.add_paragraph()
            act_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            act_para.paragraph_format.space_before = Pt(24)
            act_run = act_para.add_run(act.upper())
            act_run.bold = True
            act_run.font.name = _SCRIPT_FONT_DOCX

        slug_para = doc.add_paragraph()
        slug_para.paragraph_format.space_before = Pt(24)
        slug_para.paragraph_format.space_after = Pt(12)
        slug_run = slug_para.add_run(_slug_line(scene))
        slug_run.bold = True
        slug_run.font.name = _SCRIPT_FONT_DOCX

        body = _scene_body(scene)
        if body:
            _add_content_paragraphs(doc, body, _SCRIPT_FONT_DOCX)

    doc.save(path)


def _docx_stage_script(data: dict, path: str) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = _SCRIPT_FONT_DOCX
    style.font.size = Pt(12)
    style.paragraph_format.space_after = Pt(0)
    style.paragraph_format.line_spacing = 1.0

    for section in doc.sections:
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    _docx_title_page(doc, data["project"]["title"], _SCRIPT_FONT_DOCX)

    if not data["scenes"]:
        doc.add_paragraph("No scenes.")
        doc.save(path)
        return

    current_act = None
    scene_num = 0
    for scene in data["scenes"]:
        act = scene.get("act", "")
        if act and act != current_act:
            current_act = act
            scene_num = 0
            act_para = doc.add_paragraph()
            act_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            act_para.paragraph_format.space_before = Pt(36)
            act_run = act_para.add_run(act.upper())
            act_run.bold = True
            act_run.font.name = _SCRIPT_FONT_DOCX

        scene_num += 1
        sc_para = doc.add_paragraph()
        sc_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        sc_para.paragraph_format.space_before = Pt(24)
        sc_run = sc_para.add_run(f"SCENE {scene_num}")
        sc_run.bold = True
        sc_run.font.name = _SCRIPT_FONT_DOCX

        body = _scene_body(scene)
        if body:
            _add_content_paragraphs(doc, body, _SCRIPT_FONT_DOCX)

    doc.save(path)


def _docx_graphic_novel(data: dict, path: str) -> None:
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = _SCRIPT_FONT_DOCX
    style.font.size = Pt(12)
    style.paragraph_format.space_after = Pt(0)
    style.paragraph_format.line_spacing = 1.0

    _docx_title_page(doc, data["project"]["title"], _SCRIPT_FONT_DOCX)

    if not data["scenes"]:
        doc.add_paragraph("No scenes.")
        doc.save(path)
        return

    page_num = 0
    for scene in data["scenes"]:
        page_num += 1
        pg_para = doc.add_paragraph()
        pg_para.paragraph_format.space_before = Pt(24)
        pg_run = pg_para.add_run(f"PAGE {page_num}")
        pg_run.bold = True
        pg_run.font.name = _SCRIPT_FONT_DOCX

        body = _scene_body(scene)
        if body:
            _add_content_paragraphs(doc, body, _SCRIPT_FONT_DOCX)

    doc.save(path)


def _add_content_paragraphs(doc, text: str, font_name: str) -> None:
    from docx.shared import Pt

    for block in text.split("\n\n"):
        block = block.strip()
        if not block:
            continue
        para = doc.add_paragraph()
        lines = block.split("\n")
        for i, line in enumerate(lines):
            if i > 0:
                para.add_run().add_break()
            run = para.add_run(line)
            run.font.name = font_name
            run.font.size = Pt(12)


# -- Fountain export ----------------------------------------------------------

def export_fountain(db: Database, project_id: int) -> str:
    data = _gather_project_data(db, project_id)
    lines: list[str] = []

    lines.append(f"Title: {data['project']['title']}")
    lines.append(f"Credit: Written by")
    lines.append(f"Author: ")
    lines.append(f"Draft date: ")
    lines.append("")
    lines.append("")

    fmt = _get_fmt(data)

    current_act = None
    for scene in data["scenes"]:
        act = scene.get("act", "")
        if act and act != current_act:
            current_act = act
            if fmt in ("series", "stage_script"):
                lines.append(f"= {act}")
                lines.append("")

        slug = _slug_line(scene)
        lines.append(f".{slug}")
        lines.append("")

        body = _scene_body(scene)
        if body:
            lines.append(body)
            lines.append("")

    return "\n".join(lines)


# -- FDX (Final Draft XML) export ---------------------------------------------

def export_fdx(db: Database, project_id: int) -> str:
    data = _gather_project_data(db, project_id)

    root = ET.Element("FinalDraft", DocumentType="Script", Template="No", Version="4")
    content = ET.SubElement(root, "Content")

    title_para = ET.SubElement(content, "Paragraph", Type="Action")
    title_text = ET.SubElement(title_para, "Text")
    title_text.text = data["project"]["title"]

    fmt = _get_fmt(data)

    current_act = None
    for scene in data["scenes"]:
        act = scene.get("act", "")
        if act and act != current_act:
            current_act = act
            if fmt in ("series", "stage_script"):
                act_para = ET.SubElement(content, "Paragraph", Type="Action")
                act_t = ET.SubElement(act_para, "Text")
                act_t.text = act.upper()

        slug = _slug_line(scene)
        slug_para = ET.SubElement(content, "Paragraph", Type="Scene Heading")
        slug_text = ET.SubElement(slug_para, "Text")
        slug_text.text = slug

        body = _scene_body(scene)
        if body:
            for block in body.split("\n\n"):
                block = block.strip()
                if not block:
                    continue
                action_para = ET.SubElement(content, "Paragraph", Type="Action")
                action_text = ET.SubElement(action_para, "Text")
                action_text.text = block.replace("\n", " ")

    tree = ET.ElementTree(root)
    buf = io.BytesIO()
    tree.write(buf, encoding="utf-8", xml_declaration=True)
    return buf.getvalue().decode("utf-8")


# -- PDF export ----------------------------------------------------------------

def export_pdf(db: Database, project_id: int, path: str) -> None:
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import (
        Paragraph,
        SimpleDocTemplate,
        Spacer,
    )

    data = _gather_project_data(db, project_id)
    fmt = _get_fmt(data)
    is_script = _is_script_format(fmt)
    font_name = "Courier" if is_script else "Times-Roman"
    font_bold = "Courier-Bold" if is_script else "Times-Bold"
    font_italic = "Courier-Oblique" if is_script else "Times-Italic"

    doc = SimpleDocTemplate(
        path,
        pagesize=letter,
        leftMargin=1.5 * inch if is_script else 1.0 * inch,
        rightMargin=1.0 * inch,
        topMargin=1.0 * inch,
        bottomMargin=1.0 * inch,
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "ScriptTitle",
        parent=styles["Title"],
        fontName=font_bold,
        fontSize=24,
        alignment=1,
        spaceAfter=36,
    )
    heading_style = ParagraphStyle(
        "SceneHeading",
        parent=styles["Normal"],
        fontName=font_bold,
        fontSize=12,
        spaceBefore=24,
        spaceAfter=12,
    )
    body_style = ParagraphStyle(
        "Body",
        parent=styles["Normal"],
        fontName=font_name,
        fontSize=12,
        leading=14 if is_script else 16,
        spaceAfter=6 if is_script else 8,
    )
    act_style = ParagraphStyle(
        "Act",
        parent=styles["Normal"],
        fontName=font_bold,
        fontSize=14,
        alignment=1,
        spaceBefore=24,
        spaceAfter=12,
    )
    chapter_style = ParagraphStyle(
        "Chapter",
        parent=styles["Normal"],
        fontName=font_bold,
        fontSize=16,
        spaceBefore=24,
        spaceAfter=12,
    )

    elements: list = []
    elements.append(Paragraph(_esc(data["project"]["title"]), title_style))
    elements.append(Spacer(1, 36))

    if not data["scenes"]:
        elements.append(Paragraph("No scenes.", body_style))
        doc.build(elements)
        return

    if fmt in ("screenplay", "series"):
        _pdf_screenplay(data, fmt, elements, heading_style, body_style, act_style)
    elif fmt == "stage_script":
        _pdf_stage_script(data, elements, heading_style, body_style, act_style)
    elif fmt == "graphic_novel":
        _pdf_graphic_novel(data, elements, heading_style, body_style)
    else:
        _pdf_novel(data, elements, chapter_style, heading_style, body_style)

    doc.build(elements)


def _esc(text: str) -> str:
    return (
        text.replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
    )


def _pdf_body_blocks(text: str, style, elements: list) -> None:
    from reportlab.platypus import Paragraph

    for block in text.split("\n\n"):
        block = block.strip()
        if not block:
            continue
        safe = _esc(block).replace("\n", "<br/>")
        elements.append(Paragraph(safe, style))


def _pdf_novel(data, elements, chapter_style, heading_style, body_style):
    from reportlab.platypus import Paragraph

    chapter_groups = _group_scenes_by_chapter(data["scenes"])
    chapter_num = 0
    for chapter_name, group_scenes in chapter_groups:
        if chapter_name:
            chapter_num += 1
            elements.append(
                Paragraph(_esc(f"Chapter {chapter_num}: {chapter_name}"), chapter_style),
            )
        for scene in group_scenes:
            elements.append(Paragraph(f"<i>{_esc(scene['title'])}</i>", heading_style))
            body = _scene_body(scene)
            if body:
                _pdf_body_blocks(body, body_style, elements)


def _pdf_screenplay(data, fmt, elements, heading_style, body_style, act_style):
    from reportlab.platypus import Paragraph

    current_act = None
    for scene in data["scenes"]:
        act = scene.get("act", "")
        if fmt == "series" and act and act != current_act:
            current_act = act
            elements.append(Paragraph(_esc(act.upper()), act_style))
        elements.append(Paragraph(_esc(_slug_line(scene)), heading_style))
        body = _scene_body(scene)
        if body:
            _pdf_body_blocks(body, body_style, elements)


def _pdf_stage_script(data, elements, heading_style, body_style, act_style):
    from reportlab.platypus import Paragraph

    current_act = None
    scene_num = 0
    for scene in data["scenes"]:
        act = scene.get("act", "")
        if act and act != current_act:
            current_act = act
            scene_num = 0
            elements.append(Paragraph(_esc(act.upper()), act_style))
        scene_num += 1
        elements.append(Paragraph(_esc(f"SCENE {scene_num}"), heading_style))
        body = _scene_body(scene)
        if body:
            _pdf_body_blocks(body, body_style, elements)


def _pdf_graphic_novel(data, elements, heading_style, body_style):
    from reportlab.platypus import Paragraph

    page_num = 0
    for scene in data["scenes"]:
        page_num += 1
        elements.append(Paragraph(_esc(f"PAGE {page_num}"), heading_style))
        body = _scene_body(scene)
        if body:
            _pdf_body_blocks(body, body_style, elements)


# -- HTML export ---------------------------------------------------------------

def export_html(db: Database, project_id: int) -> str:
    data = _gather_project_data(db, project_id)
    fmt = _get_fmt(data)
    is_script = _is_script_format(fmt)
    font = "Courier New, Courier, monospace" if is_script else "Times New Roman, Georgia, serif"
    title = _esc(data["project"]["title"])

    css = (
        "body { max-width: 720px; margin: 40px auto; padding: 0 20px; "
        f"font-family: {font}; font-size: 12pt; line-height: 1.5; }}\n"
        "h1 { text-align: center; }\n"
        "h2 { text-align: center; text-transform: uppercase; font-size: 14pt; }\n"
        ".scene-heading { font-weight: bold; text-transform: uppercase; "
        "margin-top: 24px; margin-bottom: 12px; }\n"
        ".chapter { font-size: 16pt; font-weight: bold; margin-top: 36px; }\n"
        ".body-text { margin-bottom: 8px; }\n"
    )

    parts: list[str] = [
        "<!DOCTYPE html>",
        "<html lang=\"en\">",
        "<head>",
        f"<meta charset=\"utf-8\"><title>{title}</title>",
        f"<style>{css}</style>",
        "</head>",
        "<body>",
        f"<h1>{title}</h1>",
    ]

    if not data["scenes"]:
        parts.append("<p>No scenes.</p>")
    elif fmt in ("screenplay", "series"):
        _html_screenplay(data, fmt, parts)
    elif fmt == "stage_script":
        _html_stage_script(data, parts)
    elif fmt == "graphic_novel":
        _html_graphic_novel(data, parts)
    else:
        _html_novel(data, parts)

    parts.append("</body>")
    parts.append("</html>")
    return "\n".join(parts)


def _html_body_blocks(text: str, parts: list[str]) -> None:
    for block in text.split("\n\n"):
        block = block.strip()
        if not block:
            continue
        safe = _esc(block).replace("\n", "<br>")
        parts.append(f"<p class=\"body-text\">{safe}</p>")


def _html_novel(data, parts):
    chapter_groups = _group_scenes_by_chapter(data["scenes"])
    chapter_num = 0
    for chapter_name, group_scenes in chapter_groups:
        if chapter_name:
            chapter_num += 1
            parts.append(f"<div class=\"chapter\">Chapter {chapter_num}: {_esc(chapter_name)}</div>")
        for scene in group_scenes:
            parts.append(f"<div class=\"scene-heading\"><em>{_esc(scene['title'])}</em></div>")
            body = _scene_body(scene)
            if body:
                _html_body_blocks(body, parts)


def _html_screenplay(data, fmt, parts):
    current_act = None
    for scene in data["scenes"]:
        act = scene.get("act", "")
        if fmt == "series" and act and act != current_act:
            current_act = act
            parts.append(f"<h2>{_esc(act.upper())}</h2>")
        parts.append(f"<div class=\"scene-heading\">{_esc(_slug_line(scene))}</div>")
        body = _scene_body(scene)
        if body:
            _html_body_blocks(body, parts)


def _html_stage_script(data, parts):
    current_act = None
    scene_num = 0
    for scene in data["scenes"]:
        act = scene.get("act", "")
        if act and act != current_act:
            current_act = act
            scene_num = 0
            parts.append(f"<h2>{_esc(act.upper())}</h2>")
        scene_num += 1
        parts.append(f"<div class=\"scene-heading\">SCENE {scene_num}</div>")
        body = _scene_body(scene)
        if body:
            _html_body_blocks(body, parts)


def _html_graphic_novel(data, parts):
    page_num = 0
    for scene in data["scenes"]:
        page_num += 1
        parts.append(f"<div class=\"scene-heading\">PAGE {page_num}</div>")
        body = _scene_body(scene)
        if body:
            _html_body_blocks(body, parts)
