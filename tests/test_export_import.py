"""Tests for format-aware export and import."""

import json

from storyplanner.db import Database
from storyplanner.export import (
    export_formatted_text,
    export_json,
    export_manuscript,
    export_screenplay,
    export_docx_manuscript,
    export_markdown,
    export_csv_scenes,
)
from storyplanner.import_data import import_json, validate_import_data


def _make_project(db, fmt="novel"):
    proj = db.create_project("Test Story", format_mode=fmt)
    db.create_scene(proj.id, "Opening", content="The dawn broke.", act="Act One", chapter="Chapter 1")
    db.create_scene(proj.id, "Climax", content="She faced her fears.", act="Act Two", chapter="Chapter 2")
    return proj


# -- JSON round-trip preserves format_mode ------------------------------------

def test_json_export_includes_format_mode():
    db = Database()
    proj = _make_project(db, "screenplay")
    raw = export_json(db, proj.id)
    data = json.loads(raw)
    assert data["project"]["format_mode"] == "screenplay"


def test_json_export_novel_format_mode():
    db = Database()
    proj = _make_project(db, "novel")
    raw = export_json(db, proj.id)
    data = json.loads(raw)
    assert data["project"]["format_mode"] == "novel"


def test_json_roundtrip_format_mode():
    db = Database()
    proj = _make_project(db, "stage_script")
    raw = export_json(db, proj.id)
    data, err = validate_import_data(raw)
    assert err == ""
    new_id = import_json(db, data)
    new_proj = db.get_project_by_id(new_id)
    assert new_proj.format_mode == "stage_script"


def test_json_roundtrip_graphic_novel():
    db = Database()
    proj = _make_project(db, "graphic_novel")
    raw = export_json(db, proj.id)
    new_id = import_json(db, json.loads(raw))
    new_proj = db.get_project_by_id(new_id)
    assert new_proj.format_mode == "graphic_novel"


def test_json_roundtrip_series():
    db = Database()
    proj = _make_project(db, "series")
    raw = export_json(db, proj.id)
    new_id = import_json(db, json.loads(raw))
    new_proj = db.get_project_by_id(new_id)
    assert new_proj.format_mode == "series"


def test_import_missing_format_mode_defaults_novel():
    db = Database()
    data = {
        "project": {"title": "Legacy"},
        "characters": [],
        "places": [],
        "notes": [],
        "scenes": [],
    }
    new_id = import_json(db, data)
    proj = db.get_project_by_id(new_id)
    assert proj.format_mode == "novel"


# -- Format-aware text export ------------------------------------------------

def test_formatted_text_novel():
    db = Database()
    proj = _make_project(db, "novel")
    text = export_formatted_text(db, proj.id)
    assert "Test Story" in text
    assert "The dawn broke." in text
    assert "Chapter 1" in text


def test_formatted_text_screenplay():
    db = Database()
    proj = _make_project(db, "screenplay")
    text = export_formatted_text(db, proj.id)
    assert "TEST STORY" in text
    assert "OPENING" in text
    assert "The dawn broke." in text


def test_formatted_text_stage_script():
    db = Database()
    proj = _make_project(db, "stage_script")
    text = export_formatted_text(db, proj.id)
    assert "TEST STORY" in text
    assert "ACT ONE" in text
    assert "SCENE" in text


def test_formatted_text_graphic_novel():
    db = Database()
    proj = _make_project(db, "graphic_novel")
    text = export_formatted_text(db, proj.id)
    assert "PAGE 1" in text
    assert "PAGE 2" in text
    assert "The dawn broke." in text


def test_formatted_text_series():
    db = Database()
    proj = _make_project(db, "series")
    text = export_formatted_text(db, proj.id)
    assert "TEST STORY" in text
    assert "ACT ONE" in text
    assert "ACT TWO" in text


# -- Legacy exports still work -----------------------------------------------

def test_export_screenplay_function():
    db = Database()
    proj = _make_project(db)
    text = export_screenplay(db, proj.id)
    assert "TEST STORY" in text


def test_export_manuscript_function():
    db = Database()
    proj = _make_project(db)
    text = export_manuscript(db, proj.id)
    assert "Test Story" in text
    assert "Chapter 1" in text


# -- DOCX format-aware -------------------------------------------------------

def test_docx_novel(tmp_path):
    db = Database()
    proj = _make_project(db, "novel")
    path = str(tmp_path / "novel.docx")
    export_docx_manuscript(db, proj.id, path)
    from docx import Document
    doc = Document(path)
    full = "\n".join(p.text for p in doc.paragraphs)
    assert "Test Story" in full
    assert "The dawn broke." in full


def test_docx_screenplay(tmp_path):
    db = Database()
    proj = _make_project(db, "screenplay")
    path = str(tmp_path / "screenplay.docx")
    export_docx_manuscript(db, proj.id, path)
    from docx import Document
    doc = Document(path)
    full = "\n".join(p.text for p in doc.paragraphs)
    assert "Test Story" in full
    assert "OPENING" in full


def test_docx_stage_script(tmp_path):
    db = Database()
    proj = _make_project(db, "stage_script")
    path = str(tmp_path / "stage.docx")
    export_docx_manuscript(db, proj.id, path)
    from docx import Document
    doc = Document(path)
    full = "\n".join(p.text for p in doc.paragraphs)
    assert "Test Story" in full
    assert "ACT ONE" in full
    assert "SCENE 1" in full


def test_docx_graphic_novel(tmp_path):
    db = Database()
    proj = _make_project(db, "graphic_novel")
    path = str(tmp_path / "graphic.docx")
    export_docx_manuscript(db, proj.id, path)
    from docx import Document
    doc = Document(path)
    full = "\n".join(p.text for p in doc.paragraphs)
    assert "PAGE 1" in full


def test_docx_series(tmp_path):
    db = Database()
    proj = _make_project(db, "series")
    path = str(tmp_path / "series.docx")
    export_docx_manuscript(db, proj.id, path)
    from docx import Document
    doc = Document(path)
    full = "\n".join(p.text for p in doc.paragraphs)
    assert "Test Story" in full
    assert "ACT ONE" in full


# -- Other exports unaffected ------------------------------------------------

def test_markdown_export():
    db = Database()
    proj = _make_project(db)
    text = export_markdown(db, proj.id)
    assert "# Test Story" in text
    assert "## Scenes" in text


def test_csv_export():
    db = Database()
    proj = _make_project(db)
    text = export_csv_scenes(db, proj.id)
    assert "Opening" in text
    assert "Climax" in text


def test_json_roundtrip_scenes_preserved():
    db = Database()
    proj = _make_project(db, "screenplay")
    raw = export_json(db, proj.id)
    data = json.loads(raw)
    assert len(data["scenes"]) == 2
    new_id = import_json(db, data)
    scenes = db.get_all_scenes(new_id)
    assert len(scenes) == 2
    assert scenes[0].content == "The dawn broke."
    assert scenes[1].content == "She faced her fears."
